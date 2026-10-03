#!/usr/bin/env python3
"""Import Quebec MAMH 2024 municipal financial data for the 19 LGPI cities.

The MAMH "Données réelles 2024" workbook contains one SimpleOccurrence row
per municipal organism and thousands of accounting-code columns. This importer
uses current-year, actual, integral/consolidated (CRIIX) fields only and keeps
the official DescriptionPoste dictionary alongside every normalized mapping.

Quebec transparency scores are intentionally out of scope: Frontier's
published 2024 LGPI edition explicitly omitted Quebec transparency scoring.
"""

from __future__ import annotations

import io
import json
import re
import unicodedata
import urllib.parse
import urllib.request
from pathlib import Path

from openpyxl import load_workbook

SOURCE_URL = "https://mamh.gouv.qc.ca/fichiersdonneesouvertes/Données-réelles-2024.xlsx"
DATASET_URL = "https://www.donneesquebec.ca/recherche/dataset/rapport-financier-des-organismes-municipaux-et-autres-documents"
OUTPUT = Path("data/2024/quebec.json")

TARGETS = {
    "blainville": "Blainville",
    "brossard": "Brossard",
    "drummondville": "Drummondville",
    "gatineau": "Gatineau",
    "granby": "Granby",
    "laval": "Laval",
    "levis": "Levis",
    "longueuil": "Longueuil",
    "montreal": "Montreal",
    "quebec": "Quebec",
    "repentigny": "Repentigny",
    "saguenay": "Saguenay",
    "saint-hyacinthe": "Saint-Hyacinthe",
    "saint-jean-sur-richelieu": "Saint-Jean-sur-Richelieu",
    "saint-jerome": "Saint-Jerome",
    "shawinigan": "Shawinigan",
    "sherbrooke": "Sherbrooke",
    "terrebonne": "Terrebonne",
    "trois-rivieres": "Trois-Rivieres",
}

DIRECT_FIELDS = {
    "financial_assets_total": "CRIIX00632",
    "holdings_council_controlled": "CRIIX00637",
    "long_term_debt": "CRIIX00645",
    "employee_future_benefit_liability": "CRIIX00646",
    "financial_liabilities_total": "CRIIX00640",
    "net_financial_assets": "CRIIX00647",
    "capital_assets": "CRIIX00649",
    "user_charges": "CRIIX00544",
    "developer_contributions": "CRIIX00557",
    "general_government_total": "CRIIX00561",
    "public_safety_total": "CRIIX00562",
    "transportation_total": "CRIIX00563",
    "environmental_services": "CRIIX00564",
    "social_services_total": "CRIIX00565",
    "planning_development": "CRIIX00566",
    "recreation_culture": "CRIIX00567",
    "utility_operations": "CRIIX00568",
    "total_expenditure": "CRIIX00560",
    "depreciation_function": "CRIIX01888",
    "depreciation_object": "CRIIX01888",
    "total_expenditure_object": "CRIIX00661",
}

SUM_FIELDS = {
    "net_taxes": (
        ["CRIIX00540", "CRIIX00541"],
        "Taxes plus compensations in lieu of taxes.",
    ),
    "government_grants_total": (
        ["CRIIX00543", "CRIIX00554"],
        "Operating and investment transfers combined. The high-level MAMH result does not split federal and provincial transfers.",
    ),
    "investment_income": (
        ["CRIIX00547", "CRIIX00548"],
        "Portfolio investment income plus other interest income.",
    ),
    "salaries_benefits": (
        ["CRIIX01884", "CRIIX01885", "CRIIX01886", "CRIIX01887"],
        "Remuneration and social charges, including Accès entreprise Québec and other categories.",
    ),
    "interest_expense": (
        ["CRIIX00674", "CRIIX00675", "CRIIX00676", "CRIIX00677"],
        "Interest and other charges on long-term debt borne by the municipality, other municipal bodies, Quebec government/entities and other third parties.",
    ),
    "grants_object": (
        ["CRIIX00681", "CRIIX00682", "CRIIX00683", "CRIIX00685", "CRIIX00686"],
        "Contributions to municipal and other organizations: shares, transfers and other contributions.",
    ),
    "goods_services_total": (
        ["CRIIX00667", "CRIIX00669", "CRIIX00670", "CRIIX00671"],
        "Mapped MAMH goods/services components visible in the integral S19 object schedule. This is retained as pending review because the legacy LGPI grouping may include additional sub-lines.",
    ),
}

EXPECTED_DESCRIPTION_TERMS = {
    "CRIIX00632": ("Situation financière", "Actifs financiers"),
    "CRIIX00637": ("Actifs financiers", "Participations"),
    "CRIIX00640": ("Situation financière", "Passifs"),
    "CRIIX00645": ("Passifs", "Dette à long terme"),
    "CRIIX00646": ("Passifs", "avantages sociaux futurs"),
    "CRIIX00647": ("Actifs financiers nets",),
    "CRIIX00649": ("Immobilisations corporelles",),
    "CRIIX00539": ("Revenus", "Fonctionnement"),
    "CRIIX00540": ("Revenus", "Taxes"),
    "CRIIX00541": ("Compensations tenant lieu de taxes",),
    "CRIIX00543": ("Revenus", "Transferts"),
    "CRIIX00544": ("Services rendus",),
    "CRIIX00547": ("Revenus de placements",),
    "CRIIX00548": ("Autres revenus d'intérêts",),
    "CRIIX00551": ("Revenus", "Investissement"),
    "CRIIX00554": ("Investissement", "Transferts"),
    "CRIIX00557": ("Contributions des promoteurs",),
    "CRIIX00560": ("Charges",),
    "CRIIX00561": ("Charges", "Administration générale"),
    "CRIIX00562": ("Charges", "Sécurité publique"),
    "CRIIX00563": ("Charges", "Transport"),
    "CRIIX00564": ("Charges", "Hygiène du milieu"),
    "CRIIX00565": ("Charges", "Santé et bien-être"),
    "CRIIX00566": ("Charges", "Aménagement"),
    "CRIIX00567": ("Charges", "Loisirs et culture"),
    "CRIIX00568": ("Charges", "Réseau d'électricité"),
    "CRIIX00661": ("Charges par objets",),
    "CRIIX01884": ("Charges par objets", "Rémunération"),
    "CRIIX01885": ("Charges par objets", "Rémunération"),
    "CRIIX01886": ("Charges sociales",),
    "CRIIX01887": ("Charges sociales",),
    "CRIIX01888": ("Amortissement", "Immobilisations corporelles"),
    "CRIIX00674": ("Frais de financement", "dette à long terme"),
    "CRIIX00675": ("Frais de financement", "dette à long terme"),
    "CRIIX00676": ("Frais de financement", "dette à long terme"),
    "CRIIX00677": ("Frais de financement", "dette à long terme"),
    "CRIIX00681": ("Contributions", "Quotes-parts"),
    "CRIIX00682": ("Contributions", "Transferts"),
    "CRIIX00683": ("Contributions", "Autres"),
    "CRIIX00685": ("Contributions", "Transferts"),
    "CRIIX00686": ("Contributions", "Autres"),
    "CRIIX00667": ("Biens et services",),
    "CRIIX00669": ("Biens et services",),
    "CRIIX00670": ("Biens et services",),
    "CRIIX00671": ("Biens et services",),
}


def norm(value: object) -> str:
    text=unicodedata.normalize("NFKD",str(value or ""))
    text="".join(ch for ch in text if not unicodedata.combining(ch))
    return re.sub(r"[^A-Z0-9]+","",text.upper())


def safe_url(url: str) -> str:
    parsed=urllib.parse.urlsplit(url)
    return urllib.parse.urlunsplit((
        parsed.scheme,
        parsed.netloc,
        urllib.parse.quote(urllib.parse.unquote(parsed.path), safe="/"),
        parsed.query,
        parsed.fragment,
    ))


def fetch(url: str) -> bytes:
    request=urllib.request.Request(
        safe_url(url),
        headers={"User-Agent":"LGPI-reconstruction/0.1 (+https://github.com/FTFNAnalytics/lgpi)"},
    )
    with urllib.request.urlopen(request,timeout=180) as response:
        payload=response.read()
    if len(payload)<1_000_000:
        raise RuntimeError(f"Quebec MAMH workbook unexpectedly small: {len(payload)} bytes")
    return payload


def number(value: object) -> int | None:
    if value is None or value == "":
        return None
    if isinstance(value,str) and value.strip().upper() in {"N/A","NA","-","—"}:
        return None
    return int(round(float(value)))


def description_map(workbook) -> dict[str,dict]:
    ws=workbook["DescriptionPoste"]
    result={}
    for row in ws.iter_rows(min_row=2,values_only=True):
        code=str(row[0] or "").strip()
        if not code:
            continue
        result[code]={
            "subsection":str(row[1] or "").strip(),
            "subject":str(row[2] or "").strip(),
            "facets":str(row[3] or "").strip(),
        }
    return result


def validate_dictionary(descriptions: dict[str,dict], codes: set[str]) -> None:
    missing=[]
    mismatched=[]
    for code in sorted(codes):
        item=descriptions.get(code)
        if item is None:
            missing.append(code)
            continue
        terms=EXPECTED_DESCRIPTION_TERMS.get(code,())
        subject=norm(item["subject"])
        if any(norm(term) not in subject for term in terms):
            mismatched.append((code,item["subject"],terms))
    if missing or mismatched:
        raise RuntimeError(
            "MAMH dictionary validation failed: "
            + f"missing={missing}; mismatched={mismatched}"
        )


def observation(slug,metric,value,code,descriptions,status=None,note=None,method="direct",codes=None):
    source_codes=codes or [code]
    labels=[descriptions[c]["subject"] for c in source_codes]
    resolved=status or ("reported" if value is not None else "not_reported")
    return {
        "municipalitySlug":slug,
        "year":2024,
        "metric":metric,
        "valueThousands":None if value is None else value/1000,
        "status":resolved,
        "sourceUrl":SOURCE_URL,
        "sourceLabel":"Québec MAMH — Rapport financier 2024, Données réelles",
        "sourceLocation":"SimpleOccurrence: " + " + ".join(source_codes),
        "sourceReportedLabel":" + ".join(labels),
        "sourceFieldCodes":source_codes,
        "mappingMethod":method,
        **({"note":note} if note else {}),
    }


def main() -> None:
    payload=fetch(SOURCE_URL)
    workbook=load_workbook(io.BytesIO(payload),read_only=True,data_only=True)
    required_sheets={"SimpleOccurrence","DescriptionPoste","OrganismeAbsent"}
    if not required_sheets.issubset(workbook.sheetnames):
        raise RuntimeError(f"Missing expected MAMH sheets: {required_sheets - set(workbook.sheetnames)}")

    descriptions=description_map(workbook)
    all_codes=set(DIRECT_FIELDS.values()) | {"CRIIX00539","CRIIX00551"}
    for codes,_ in SUM_FIELDS.values():
        all_codes.update(codes)
    validate_dictionary(descriptions,all_codes)

    ws=workbook["SimpleOccurrence"]
    headers=[str(v) if v is not None else "" for v in next(ws.iter_rows(min_row=1,max_row=1,values_only=True))]
    idx={name:i for i,name in enumerate(headers)}
    missing_columns=[code for code in all_codes if code not in idx]
    if missing_columns:
        raise RuntimeError(f"Required MAMH accounting columns absent: {missing_columns}")

    target_norms={norm(name):slug for slug,name in TARGETS.items()}
    rows={}
    for values in ws.iter_rows(min_row=2,values_only=True):
        source_name=values[idx["nom_organisme"]]
        slug=target_norms.get(norm(source_name))
        if slug:
            if slug in rows:
                raise RuntimeError(f"Duplicate MAMH target row: {slug}")
            rows[slug]=values

    if set(rows)!=set(TARGETS):
        raise RuntimeError(f"Quebec target coverage mismatch: missing={sorted(set(TARGETS)-set(rows))}")

    municipalities=[]
    observations=[]
    for slug in TARGETS:
        row=rows[slug]
        municipalities.append({
            "slug":slug,
            "name":str(row[idx["nom_organisme"]]),
            "province":"QC",
            "sourceCode":str(row[idx["cod_geo"]]),
            "sourceDesignation":row[idx["desi_org"]],
            "population":number(row[idx["population"]]),
            "transparencyStatus":"not_published_in_2024_lgpi",
        })

        values={code:number(row[idx[code]]) for code in all_codes}

        # Direct fields.
        for metric,code in DIRECT_FIELDS.items():
            note=None
            status=None
            if metric=="holdings_council_controlled":
                note="MAMH source is participations in municipal enterprises and commercial partnerships; retained as the closest legacy LGPI council-controlled-operations field."
                status="pending_review"
            observations.append(observation(slug,metric,values[code],code,descriptions,status=status,note=note))

        # Derived sums.
        for metric,(codes,note) in SUM_FIELDS.items():
            parts=[values[c] for c in codes]
            amount=None if any(v is None for v in parts) else sum(v for v in parts if v is not None)
            status="pending_review" if metric=="goods_services_total" else None
            observations.append(observation(
                slug,metric,amount,codes[0],descriptions,status=status,note=note,
                method="derived_sum",codes=codes,
            ))

        # Total revenue = operating + investment revenue.
        revenue_codes=["CRIIX00539","CRIIX00551"]
        parts=[values[c] for c in revenue_codes]
        total_revenue=None if any(v is None for v in parts) else sum(v for v in parts if v is not None)
        observations.append(observation(
            slug,"total_revenue",total_revenue,revenue_codes[0],descriptions,
            note="Operating plus investment revenue from the integral current-year actual result.",
            method="derived_sum",codes=revenue_codes,
        ))

        # Other revenue as an exhaustive residual of the high-level mapped categories.
        mapped_metrics={"net_taxes","government_grants_total","user_charges","investment_income","developer_contributions"}
        current={item["metric"]:item for item in observations if item["municipalitySlug"]==slug}
        mapped_values=[current[m]["valueThousands"] for m in mapped_metrics]
        if total_revenue is None or any(v is None for v in mapped_values):
            other=None
        else:
            other=total_revenue - int(round(sum(float(v)*1000 for v in mapped_values if v is not None)))
        observations.append(observation(
            slug,"revenue_other",other,revenue_codes[0],descriptions,
            note="Total revenue less mapped taxes, transfers, services rendered, investment/interest income and developer contributions. This residual preserves rights, fines, quota shares and other revenue without dropping them.",
            method="derived_residual",codes=revenue_codes,
        ))

        # Financial-assets residual excluding municipal-enterprise/partnership participation.
        fa=values["CRIIX00632"]
        holdings=values["CRIIX00637"]
        other_fa=None if fa is None or holdings is None else fa-holdings
        observations.append(observation(
            slug,"financial_assets_other",other_fa,"CRIIX00632",descriptions,
            note="Total financial assets less participations in municipal enterprises and commercial partnerships.",
            method="derived_residual",codes=["CRIIX00632","CRIIX00637"],
        ))

        # Internal accounting controls.
        total_function=values["CRIIX00560"]
        total_object=values["CRIIX00661"]
        if total_function is not None and total_object is not None and total_function != total_object:
            raise RuntimeError(
                f"{slug}: MAMH total charges mismatch S12={total_function} S19={total_object}"
            )
        net=values["CRIIX00647"]
        assets=values["CRIIX00632"]
        liabilities=values["CRIIX00640"]
        if None not in (net,assets,liabilities) and net != assets-liabilities:
            raise RuntimeError(
                f"{slug}: net financial assets mismatch: {net} != {assets} - {liabilities}"
            )

    document={
        "schemaVersion":1,
        "province":"QC",
        "year":2024,
        "source":{
            "publisher":"Ministère des Affaires municipales et de l'Habitation (Québec)",
            "title":"Rapport financier 2024 — Données réelles",
            "datasetUrl":DATASET_URL,
            "resourceUrl":SOURCE_URL,
            "format":"XLSX",
        },
        "mapping":{
            "name":"LGPI Quebec MAMH 2024",
            "version":1,
            "notes":[
                "Only current-year actual integral/consolidated CRIIX fields are used.",
                "Values are divided by 1,000 for the legacy LGPI thousands-of-dollars display convention.",
                "Quebec transparency scores were not published in the 2024 LGPI edition and are not calculated by this importer.",
                "Federal/provincial transfer splits and detailed police/fire/transit sub-functions remain unmapped rather than inferred.",
                "Total charges must reconcile between MAMH S12 function reporting and S19 object reporting.",
                "Net financial assets must reconcile to financial assets less liabilities.",
            ],
        },
        "municipalities":municipalities,
        "observations":observations,
        "stats":{
            "municipalityCount":len(municipalities),
            "observationCount":len(observations),
            "reportedCount":sum(1 for x in observations if x["status"]=="reported"),
            "pendingReviewCount":sum(1 for x in observations if x["status"]=="pending_review"),
            "notReportedCount":sum(1 for x in observations if x["status"]=="not_reported"),
        },
    }

    OUTPUT.parent.mkdir(parents=True,exist_ok=True)
    OUTPUT.write_text(json.dumps(document,indent=2,ensure_ascii=False)+"\n",encoding="utf-8")
    print(f"Wrote {OUTPUT}: {len(municipalities)} municipalities, {len(observations)} observations")
    for m in municipalities:
        print(f"- {m['slug']}: {m['name']!r} code={m['sourceCode']} population={m['population']}")


if __name__=="__main__":
    main()

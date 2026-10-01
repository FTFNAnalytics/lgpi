#!/usr/bin/env python3
"""Inspect Alberta's official 2023 municipal FIR/SIR workbook.

This is a research helper used to learn the workbook schema before the
production importer is locked. It does not write application data.
"""

from __future__ import annotations

import io
import urllib.request

from openpyxl import load_workbook

SOURCE_URL = "https://open.alberta.ca/dataset/cde4c4fd-a0b2-4816-af43-13de7a3fd3e3/resource/78ee285e-460a-44d9-898a-a50b30bb1341/download/2023_financial_year.xlsx"


def download() -> bytes:
    req = urllib.request.Request(
        SOURCE_URL,
        headers={"User-Agent": "LGPI-reconstruction/0.1 (+https://github.com/FTFNAnalytics/lgpi)"},
    )
    with urllib.request.urlopen(req, timeout=60) as response:
        payload = response.read()
    print(f"downloaded_bytes={len(payload)}")
    return payload


def compact(value):
    if value is None:
        return None
    text = str(value).strip()
    return text[:180]


def main() -> None:
    payload = download()
    workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    print("sheet_names=", workbook.sheetnames)

    for name in workbook.sheetnames:
        ws = workbook[name]
        print(f"\n=== SHEET {name!r} rows={ws.max_row} cols={ws.max_column} ===")
        for idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            values = [compact(value) for value in row]
            if any(value not in (None, "") for value in values):
                print(f"ROW {idx}: {values}")
            if idx >= 18:
                break


if __name__ == "__main__":
    main()

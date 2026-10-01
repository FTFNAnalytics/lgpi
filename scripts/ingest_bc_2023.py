#!/usr/bin/env python3
"""Import British Columbia's official 2023 municipal financial schedules.

The Province of British Columbia publishes standardized municipal financial
statistics by schedule. This importer maps the strongest directly comparable
2023 fields into the legacy LGPI metric catalogue while retaining the original
schedule/column labels for provenance.

No per-household calculations are produced here.
"""

from __future__ import annotations

import io
import json
import re
import urllib.request
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

BASE = "https://www2.gov.bc.ca/assets/gov/british-columbians-our-governments/local-governments/finance/local-government-statistics"
INDEX_URL = "https://www2.gov.bc.ca/gov/content/governments/local-governments/facts-framework/statistics/statistics"
OUTPUT = Path("data/2023/bc.json")

SCHEDULE_URLS = {
    "201": f"{BASE}/schedule201_2023.xlsx",
    "301": f"{BASE}/schedule301_2023.xlsx",
    "302": f"{BASE}/schedule302_2023.xlsx",
    "304": f"{BASE}/schedule304_2023.xlsx",
    "401": f"{BASE}/schedule401_2023.xlsx",
    "402": f"{BASE}/schedule402_2023.xlsx",
    "502": f"{BASE}/schedule502_2023.xlsx",
    "601.1": f"{BASE}/schedule601_1_2023.xlsx",
}

TARGETS = {
    "abbotsford": ("ABBOTSFORD",),
    "burnaby": ("BURNABY",),
    "chilliwack": ("CHILLIWACK",),
    "coquitlam": ("COQUITLAM",),
    "delta": ("DELTA",),
    "kamloops": ("KAMLOOPS",),
    "kelowna": ("KELOWNA",),
    "langley-township": ("LANGLEY",),
    "maple-ridge": ("MAPLE RIDGE",),
    "nanaimo": ("NANAIMO",),
    "new-westminster": ("NEW WESTMINSTER",),
    "port-coquitlam": ("PORT COQUITLAM",),
    "prince-george": ("PRINCE GEORGE",),
    "richmond": ("RICHMOND",),
    "saanich": ("SAANICH",),
    "vancouver": ("VANCOUVER",),
    "victoria": ("VICTORIA",),
}

TARGET_TYPES = {
    "langley-township": "D",
}


def fetch(url: str) -> bytes:
    request = urllib.request.Request(
        url,
        headers={"User-Agent": "LGPI-reconstruction/0.1 (+https://github.com/FTFNAnalytics/lgpi)"},
    )
    with urllib.request.urlopen(request, timeout=90) as response:
        payload = response.read()
    if len(payload) < 10_000:
        raise RuntimeError(f"BC source download unexpectedly small: {url} ({len(payload)} bytes)")
    return payload


def normalize_name(value: object) -> str:
    return re.sub(r"[^A-Z0-9]+", "", str(value or "").upper())


def numeric(value: object) -> int | None:
    if value is None or value == "":
        return None
    if isinstance(value, str) and value.strip().upper() in {"N/A", "NA", "-", "—"}:
        return None
    return int(round(float(value)))


@dataclass(frozen=True)
class Schedule:
    code: str
    url: str
    headers: list[str]
    rows: dict[str, dict[str, object]]


def load_schedule(code: str, url: str) -> Schedule:
    workbook = load_workbook(io.BytesIO(fetch(url)), read_only=True, data_only=True)
    ws = workbook[workbook.sheetnames[0]]
    headers = [str(v).strip() if v is not None else "" for v in next(ws.iter_rows(min_row=2, max_row=2, values_only=True))]
    rows: dict[str, dict[str, object]] = {}
    for row_number, values in enumerate(ws.iter_rows(min_row=3, values_only=True), start=3):
        if not values or values[0] is None:
            continue
        name = str(values[0]).strip()
        row = {header: values[i] if i < len(values) else None for i, header in enumerate(headers) if header}
        row["_source_row"] = row_number
        row_key = f"{name}::{row.get('Type')}::{row.get('RD')}::{row_number}"
        rows[row_key] = row
    return Schedule(code=code, url=url, headers=headers, rows=rows)


def find_target_names(schedule: Schedule) -> dict[str, str]:
    found: dict[str, str] = {}
    for slug, aliases in TARGETS.items():
        needles = {normalize_name(alias) for alias in aliases}
        expected_type = TARGET_TYPES.get(slug)

        def eligible(item: tuple[str, dict[str, object]]) -> bool:
            _, row = item
            if expected_type is None:
                return True
            return str(row.get("Type") or "").strip().upper() == expected_type

        exact = [
            key for key, row in schedule.rows.items()
            if eligible((key, row))
            and normalize_name(row.get("Municipalities")) in needles
        ]
        fuzzy = [
            key for key, row in schedule.rows.items()
            if eligible((key, row))
            and any(needle in normalize_name(row.get("Municipalities")) for needle in needles)
        ]
        matches = exact if exact else fuzzy
        if len(matches) != 1:
            detail = [
                {
                    "name": schedule.rows[key].get("Municipalities"),
                    "type": schedule.rows[key].get("Type"),
                    "rd": schedule.rows[key].get("RD"),
                    "source_row": schedule.rows[key].get("_source_row"),
                }
                for key in matches
            ]
            raise RuntimeError(
                f"Expected one BC source row for {slug}; "
                f"expected_type={expected_type!r}, matches={detail}"
            )
        found[slug] = matches[0]
    return found


def value(schedule: Schedule, row_name: str, label: str) -> int | None:
    if label not in schedule.headers:
        raise RuntimeError(f"{schedule.code}: missing expected column {label!r}")
    return numeric(schedule.rows[row_name].get(label))


def observation(
    *,
    slug: str,
    metric: str,
    amount: int | None,
    schedule: Schedule,
    labels: list[str],
    method: str,
    note: str | None = None,
    status: str | None = None,
) -> dict:
    resolved_status = status or ("reported" if amount is not None else "not_reported")
    return {
        "municipalitySlug": slug,
        "year": 2023,
        "metric": metric,
        "valueThousands": None if amount is None else amount / 1000,
        "status": resolved_status,
        "sourceUrl": schedule.url,
        "sourceLabel": f"Province of British Columbia — Municipal Financial Statistics 2023, Schedule {schedule.code}",
        "sourceLocation": f"Schedule {schedule.code}: " + " + ".join(labels),
        "sourceReportedLabel": " + ".join(labels),
        "sourceFieldCodes": [f"{schedule.code}:{label}" for label in labels],
        "mappingMethod": method,
        **({"note": note} if note else {}),
    }


def direct(slug: str, metric: str, schedule: Schedule, row_name: str, label: str, *, note: str | None = None, status: str | None = None) -> dict:
    return observation(
        slug=slug,
        metric=metric,
        amount=value(schedule, row_name, label),
        schedule=schedule,
        labels=[label],
        method="direct",
        note=note,
        status=status,
    )


def derived_sum(slug: str, metric: str, schedule: Schedule, row_name: str, labels: list[str], note: str) -> dict:
    values = [value(schedule, row_name, label) for label in labels]
    amount = None if any(item is None for item in values) else sum(item for item in values if item is not None)
    return observation(
        slug=slug,
        metric=metric,
        amount=amount,
        schedule=schedule,
        labels=labels,
        method="derived_sum",
        note=note,
    )


def derived_residual(slug: str, metric: str, schedule: Schedule, row_name: str, total_label: str, subtract_labels: list[str], note: str) -> dict:
    total = value(schedule, row_name, total_label)
    parts = [value(schedule, row_name, label) for label in subtract_labels]
    amount = None if total is None or any(item is None for item in parts) else total - sum(item for item in parts if item is not None)
    return observation(
        slug=slug,
        metric=metric,
        amount=amount,
        schedule=schedule,
        labels=[total_label, *subtract_labels],
        method="derived_residual",
        note=note,
    )


def validate_row(slug: str, row_names: dict[str, str], schedules: dict[str, Schedule]) -> None:
    r301 = row_names["301"]
    r302 = row_names["302"]
    r304 = row_names["304"]
    r502 = row_names["502"]
    r601 = row_names["601.1"]

    assets_301 = value(schedules["301"], r301, "Total Financial Assets")
    assets_304 = value(schedules["304"], r304, "Total Financial Assets")
    liabilities_302 = value(schedules["302"], r302, "Total Liabilities")
    liabilities_304 = value(schedules["304"], r304, "Total Liabilities")
    net_304 = value(schedules["304"], r304, "Net Financial Assets (Debt)")
    nfa_502 = value(schedules["502"], r502, "Total Non-Financial Assets")
    nfa_304 = value(schedules["304"], r304, "Total Non-Financial Assets")
    debt_302 = value(schedules["302"], r302, "Long-Term Debt")
    debt_601 = value(schedules["601.1"], r601, "Total Debt at Year End")

    failures = []
    if assets_301 != assets_304:
        failures.append(f"financial assets 301={assets_301} 304={assets_304}")
    if liabilities_302 != liabilities_304:
        failures.append(f"liabilities 302={liabilities_302} 304={liabilities_304}")
    if assets_304 is not None and liabilities_304 is not None and net_304 != assets_304 - liabilities_304:
        failures.append(f"net financial assets expected={assets_304 - liabilities_304} got={net_304}")
    if nfa_502 != nfa_304:
        failures.append(f"non-financial assets 502={nfa_502} 304={nfa_304}")
    if debt_302 != debt_601:
        failures.append(f"long-term debt 302={debt_302} 601.1={debt_601}")
    if failures:
        raise RuntimeError(f"BC cross-schedule reconciliation failed for {slug}: " + "; ".join(failures))


def build_observations(slug: str, row_names: dict[str, str], schedules: dict[str, Schedule]) -> list[dict]:
    s301, s302, s304 = schedules["301"], schedules["302"], schedules["304"]
    s401, s402, s502, s601 = schedules["401"], schedules["402"], schedules["502"], schedules["601.1"]
    r301, r302, r304 = row_names["301"], row_names["302"], row_names["304"]
    r401, r402, r502, r601 = row_names["401"], row_names["402"], row_names["502"], row_names["601.1"]

    out = [
        direct(slug, "financial_assets_total", s301, r301, "Total Financial Assets"),
        direct(slug, "holdings_council_controlled", s301, r301, "Government Business Enterprise Equity"),
        derived_residual(
            slug, "financial_assets_other", s301, r301, "Total Financial Assets",
            ["Government Business Enterprise Equity"],
            "All financial assets other than Government Business Enterprise equity.",
        ),
        direct(slug, "long_term_debt", s302, r302, "Long-Term Debt"),
        direct(
            slug, "employee_future_benefit_liability", s302, r302, "Future Obligations",
            note="BC's source label is broader than the legacy LGPI 'employee future benefit liability' field; retained for human review.",
            status="pending_review",
        ),
        direct(slug, "financial_liabilities_total", s302, r302, "Total Liabilities"),
        direct(slug, "net_financial_assets", s304, r304, "Net Financial Assets (Debt)"),
        direct(slug, "capital_assets", s502, r502, "Total Tangible Capital Assets"),

        direct(slug, "net_taxes", s401, r401, "Total Own Purpose Taxation and Grants in Lieu"),
        direct(slug, "user_charges", s401, r401, "Sale of Services"),
        direct(slug, "federal_grants", s401, r401, "Federal Government Transfers"),
        direct(slug, "provincial_grants", s401, r401, "Provincial Government Transfers"),
        derived_sum(
            slug, "government_grants_total", s401, r401,
            ["Federal Government Transfers", "Provincial Government Transfers", "Regional and Other Governments Transfers"],
            "Federal, provincial, regional and other-government transfers.",
        ),
        direct(slug, "investment_income", s401, r401, "Investment Income"),
        direct(slug, "developer_contributions", s401, r401, "Developer and Other Contributions/ Donations"),
        derived_sum(
            slug, "revenue_other", s401, r401,
            ["Income from Government Business Enterprise", "Gain on Sale of Assets", "Other Revenue"],
            "Government-business-enterprise income, gains on asset sales and other revenue.",
        ),
        direct(slug, "total_revenue", s401, r401, "Total Revenue"),

        direct(slug, "general_government_total", s402, r402, "General Government"),
        direct(slug, "public_safety_total", s402, r402, "Protective Services"),
        direct(slug, "solid_waste", s402, r402, "Solid Waste Management and Recycling"),
        direct(slug, "social_services_total", s402, r402, "Health, Social Services and Housing"),
        direct(slug, "planning_development", s402, r402, "Development Services"),
        direct(slug, "transportation_total", s402, r402, "Transportation and Transit"),
        direct(slug, "recreation_culture", s402, r402, "Parks, Recreation and Culture"),
        derived_sum(
            slug, "environmental_services", s402, r402,
            ["Solid Waste Management and Recycling", "Water Services", "Sewer Services"],
            "Solid waste, water and sewer services.",
        ),
        derived_sum(
            slug, "expenditure_other", s402, r402,
            ["Other Services", "Asset Retirement Obligation Accretion", "Loss on Disposition of Assets", "Other Adjustments"],
            "Other services plus source-specified adjustments outside the principal functional categories.",
        ),
        direct(slug, "depreciation_function", s402, r402, "Amortization"),
        direct(slug, "total_expenditure", s402, r402, "Total Expenses"),

        direct(slug, "interest_expense", s601, r601, "Interest Expense"),
        direct(
            slug, "depreciation_object", s402, r402, "Amortization",
            note="BC publishes amortization in its consolidated expense schedule; the same total is surfaced in the legacy LGPI depreciation-by-object field.",
        ),
    ]
    return out


def main() -> None:
    schedules = {code: load_schedule(code, url) for code, url in SCHEDULE_URLS.items()}

    base_names = find_target_names(schedules["301"])
    names_by_schedule: dict[str, dict[str, str]] = {"301": base_names}
    for code, schedule in schedules.items():
        if code == "301":
            continue
        names_by_schedule[code] = find_target_names(schedule)

    municipalities = []
    observations: list[dict] = []

    for slug in TARGETS:
        row_names = {code: names_by_schedule[code][slug] for code in schedules}
        validate_row(slug, row_names, schedules)
        observations.extend(build_observations(slug, row_names, schedules))

        stats_row = schedules["201"].rows[row_names["201"]]
        municipalities.append({
            "slug": slug,
            "sourceMunicipality": schedules["301"].rows[row_names["301"]].get("Municipalities"),
            "sourceType": schedules["301"].rows[row_names["301"]].get("Type"),
            "regionalDistrictCode": schedules["301"].rows[row_names["301"]].get("RD"),
            "population2021Census": numeric(stats_row.get("2021 Census")),
            "bcStatsPopulationEstimate": numeric(stats_row.get("BC Stats Population Estimates")),
        })

    document = {
        "schemaVersion": 1,
        "province": "BC",
        "year": 2023,
        "source": {
            "publisher": "Province of British Columbia",
            "title": "Municipal general and financial statistics — 2023 schedules",
            "indexUrl": INDEX_URL,
            "scheduleUrls": SCHEDULE_URLS,
        },
        "mapping": {
            "name": "LGPI British Columbia schedules 2023",
            "version": 1,
            "notes": [
                "Source workbooks report dollars; output values are divided by 1,000 to match the legacy LGPI display convention.",
                "Each LGPI municipality must reconcile across schedules 301, 302, 304, 502 and 601.1 before data is emitted.",
                "BC Future Obligations is retained as pending_review for the narrower legacy employee-future-benefit field.",
                "No population or provincial schedule statistic is used as an LGPI per-household denominator.",
            ],
        },
        "municipalities": municipalities,
        "observations": observations,
        "stats": {
            "municipalityCount": len(municipalities),
            "observationCount": len(observations),
            "reportedCount": sum(1 for row in observations if row["status"] == "reported"),
            "pendingReviewCount": sum(1 for row in observations if row["status"] == "pending_review"),
        },
    }

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(document, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(f"Wrote {OUTPUT}: {len(municipalities)} municipalities, {len(observations)} observations")
    for row in municipalities:
        print(
            f"- {row['slug']}: source={row['sourceMunicipality']!r} "
            f"type={row['sourceType']} rd={row['regionalDistrictCode']} "
            f"census={row['population2021Census']} bc_est={row['bcStatsPopulationEstimate']}"
        )


if __name__ == "__main__":
    main()

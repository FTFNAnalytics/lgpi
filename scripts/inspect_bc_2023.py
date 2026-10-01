#!/usr/bin/env python3
"""Inspect British Columbia's official 2023 municipal financial schedules."""

from __future__ import annotations

import io
import re
import urllib.parse
import urllib.request

from bs4 import BeautifulSoup
from openpyxl import load_workbook

INDEX_URL = "https://www2.gov.bc.ca/gov/content/governments/local-governments/facts-framework/statistics/statistics"


def fetch(url: str) -> bytes:
    request = urllib.request.Request(
        url,
        headers={"User-Agent": "LGPI-reconstruction/0.1 (+https://github.com/FTFNAnalytics/lgpi)"},
    )
    with urllib.request.urlopen(request, timeout=60) as response:
        return response.read()


def compact(value):
    if value is None:
        return None
    return str(value).strip()[:180]


def main() -> None:
    html = fetch(INDEX_URL)
    soup = BeautifulSoup(html, "html.parser")
    links: list[tuple[str, str]] = []
    for anchor in soup.find_all("a", href=True):
        href = urllib.parse.urljoin(INDEX_URL, anchor["href"])
        text = " ".join(anchor.stripped_strings)
        if "2023" in text and re.search(r"\.xlsx(?:$|\?)", href, re.I):
            links.append((text, href))
        elif "2023" in href and re.search(r"\.xlsx(?:$|\?)", href, re.I):
            links.append((text, href))

    # De-duplicate while preserving source order.
    seen = set()
    deduped = []
    for item in links:
        if item[1] not in seen:
            seen.add(item[1])
            deduped.append(item)

    print(f"index_bytes={len(html)} 2023_xlsx_links={len(deduped)}")
    for idx, (text, href) in enumerate(deduped, start=1):
        print(f"LINK {idx}: text={text!r} url={href}")

    # Inspect every 2023 workbook whose filename suggests one of the core LGPI schedules.
    wanted_tokens = ("201", "301", "302", "304", "401", "402", "502", "503", "601", "706")
    for text, href in deduped:
        low = href.lower()
        if not any(token in low for token in wanted_tokens):
            continue
        payload = fetch(href)
        workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
        print(f"\n=== WORKBOOK {href} bytes={len(payload)} sheets={workbook.sheetnames} ===")
        for sheet_name in workbook.sheetnames:
            ws = workbook[sheet_name]
            print(f"-- SHEET {sheet_name!r} rows={ws.max_row} cols={ws.max_column}")
            for idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
                values = [compact(value) for value in row]
                if any(value not in (None, "") for value in values):
                    print(f"ROW {idx}: {values}")
                if idx >= 10:
                    break

            if "schedule301_2023" in href.lower():
                print("-- TARGET NAME PROBE --")
                for row in ws.iter_rows(min_row=3, values_only=True):
                    name = str(row[0] or "")
                    normalized = re.sub(r"[^A-Z0-9]+", "", name.upper())
                    probes = (
                        "ABBOTSFORD","BURNABY","CHILLIWACK","COQUITLAM","DELTA",
                        "KAMLOOPS","KELOWNA","LANGLEY","MAPLERIDGE","NANAIMO",
                        "NEWWESTMINSTER","PORTCOQUITLAM","PRINCEGEORGE","RICHMOND",
                        "SAANICH","VANCOUVER","VICTORIA"
                    )
                    if any(probe in normalized for probe in probes):
                        print(f"TARGET: {name!r} type={row[1]!r} rd={row[2]!r}")


if __name__ == "__main__":
    main()

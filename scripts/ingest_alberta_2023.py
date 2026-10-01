#!/usr/bin/env python3
"""Import 2023 Alberta FIR/SIR data into the LGPI reconstruction.

Source: Government of Alberta, Municipal Financial and Statistical Data,
2023 financial year. The workbook is a standardized annual return supplied
by Alberta municipalities.

The importer only emits LGPI municipalities already present in the 2023
transparency universe. It preserves source labels/codes and marks whether a
normalized field is a direct source field or a documented sum/residual.
"""

from __future__ import annotations

import io
import json
import re
import urllib.request
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

SOURCE_URL = "https://open.alberta.ca/dataset/cde4c4fd-a0b2-4816-af43-13de7a3fd3e3/resource/78ee285e-460a-44d9-898a-a50b30bb1341/download/2023_financial_year.xlsx"
SOURCE_DATASET_URL = "https://open.alberta.ca/opendata/cde4c4fd-a0b2-4816-af43-13de7a3fd3e3"
OUTPUT = Path("data/2023/alberta.json")

TARGETS = {
    "calgary": ("CALGARY",),
    "edmonton": ("EDMONTON",),
    "grande-prairie": ("GRANDE PRAIRIE",),
    "lethbridge": ("LETHBRIDGE",),
    "medicine-hat": ("MEDICINE HAT",),
    "red-deer": ("RED DEER",),
    "st-albert": ("ST ALBERT",),
    "strathcona-county": ("STRATHCONA COUNTY",),
    "wood-buffalo": ("WOOD BUFFALO",),
}

# Independent controls transcribed from Calgary's audited 2023 report.
# These are dollars and are intentionally a small set of high-value checks.
CALGARY_CONTROLS = {
    "financial_assets_total": 10_621_576_000,
    "financial_liabilities_total": 6_098_275_000,
    "net_financial_assets": 4_523_301_000,
    "capital_assets": 20_319_980_000,
    "long_term_debt": 2_700_337_000,
    "net_taxes": 2_607_604_000,
    "user_charges": 1_359_983_000,
    "investment_income": 219_934_000,
    "developer_contributions": 477_525_000,
    "government_grants_total": 641_622_000,
    "total_revenue": 5_546_036_000,
    "total_expenditure": 4_657_747_000,
    "salaries_benefits": 2_237_853_000,
    "contracted_services": 597_112_000,
    "goods": 718_577_000,
    "grants_object": 238_615_000,
    "interest_expense": 116_885_000,
    "depreciation_object": 724_479_000,
    "total_expenditure_object": 4_657_747_000,
}


def normalize_name(value: object) -> str:
    text = str(value or "").upper()
    return re.sub(r"[^A-Z0-9]+", "", text)


def download() -> bytes:
    request = urllib.request.Request(
        SOURCE_URL,
        headers={"User-Agent": "LGPI-reconstruction/0.1 (+https://github.com/FTFNAnalytics/lgpi)"},
    )
    with urllib.request.urlopen(request, timeout=90) as response:
        payload = response.read()
    if len(payload) < 100_000:
        raise RuntimeError(f"Workbook download unexpectedly small: {len(payload)} bytes")
    return payload


@dataclass(frozen=True)
class Sheet:
    name: str
    headers: dict[str, str]
    rows: dict[str, dict[str, object]]


def load_sheet(workbook, name: str) -> Sheet:
    ws = workbook[name]
    labels = list(next(ws.iter_rows(min_row=2, max_row=2, values_only=True)))
    codes = list(next(ws.iter_rows(min_row=3, max_row=3, values_only=True)))
    headers: dict[str, str] = {}
    for label, code in zip(labels, codes):
        if code is not None:
            headers[str(code).strip()] = str(label).strip() if label is not None else str(code).strip()

    rows: dict[str, dict[str, object]] = {}
    for values in ws.iter_rows(min_row=4, values_only=True):
        if not values or values[2] is None:
            continue
        municipality_code = str(values[2]).strip().zfill(4)
        row: dict[str, object] = {
            "_year": values[0],
            "_status": values[1],
            "_code": municipality_code,
            "_municipality": values[3],
        }
        for idx, source_code in enumerate(codes):
            if source_code is None or idx >= len(values):
                continue
            row[str(source_code).strip()] = values[idx]
        rows[municipality_code] = row
    return Sheet(name=name, headers=headers, rows=rows)


def number(value: object) -> int:
    if value in (None, ""):
        return 0
    return int(round(float(value)))


def sum_codes(sheet: Sheet, row: dict[str, object], codes: Iterable[str]) -> int:
    return sum(number(row.get(code)) for code in codes)


def observation(
    *,
    slug: str,
    metric: str,
    value_dollars: int,
    sheet: Sheet,
    source_codes: list[str],
    method: str,
    note: str | None = None,
) -> dict:
    labels = [sheet.headers.get(code, code) for code in source_codes]
    return {
        "municipalitySlug": slug,
        "year": 2023,
        "metric": metric,
        "valueThousands": value_dollars / 1000,
        "status": "reported",
        "sourceUrl": SOURCE_URL,
        "sourceLabel": "Government of Alberta — Municipal Financial and Statistical Data, 2023 financial year",
        "sourceLocation": f"{sheet.name}: " + " + ".join(f"{label} [{code}]" for label, code in zip(labels, source_codes)),
        "sourceReportedLabel": " + ".join(labels),
        "sourceFieldCodes": source_codes,
        "mappingMethod": method,
        **({"note": note} if note else {}),
    }


def direct(slug: str, metric: str, sheet: Sheet, row: dict[str, object], code: str, note: str | None = None) -> dict:
    return observation(
        slug=slug,
        metric=metric,
        value_dollars=number(row.get(code)),
        sheet=sheet,
        source_codes=[code],
        method="direct",
        note=note,
    )


def derived_sum(slug: str, metric: str, sheet: Sheet, row: dict[str, object], codes: list[str], note: str) -> dict:
    return observation(
        slug=slug,
        metric=metric,
        value_dollars=sum_codes(sheet, row, codes),
        sheet=sheet,
        source_codes=codes,
        method="derived_sum",
        note=note,
    )


def find_target_codes(sheet_a: Sheet) -> dict[str, str]:
    found: dict[str, str] = {}
    for slug, needles in TARGETS.items():
        normalized_needles = [normalize_name(needle) for needle in needles]

        exact_matches = []
        fuzzy_matches = []
        for code, row in sheet_a.rows.items():
            source_name = row.get("_municipality")
            normalized = normalize_name(source_name)
            if normalized in normalized_needles:
                exact_matches.append((code, source_name))
            elif any(needle in normalized for needle in normalized_needles):
                fuzzy_matches.append((code, source_name))

        matches = exact_matches if exact_matches else fuzzy_matches
        if len(matches) != 1:
            raise RuntimeError(
                f"Expected exactly one Alberta source match for {slug}; "
                f"exact={exact_matches}, fuzzy={fuzzy_matches}"
            )
        found[slug] = matches[0][0]
    return found


def build_observations(slug: str, code: str, sheets: dict[str, Sheet]) -> list[dict]:
    a = sheets["A(1)-Total"]
    c = sheets["C(1)-Revenue"]
    d = sheets["D(1)-Total"]
    row_a, row_c, row_d = a.rows[code], c.rows[code], d.rows[code]

    out: list[dict] = [
        direct(slug, "financial_assets_total", a, row_a, "00260"),
        direct(slug, "long_term_debt", a, row_a, "00350"),
        direct(slug, "financial_liabilities_total", a, row_a, "00390"),
        direct(slug, "net_financial_assets", a, row_a, "00395"),
        direct(slug, "capital_assets", a, row_a, "00400"),

        # Revenue: these groupings reconcile exactly to Calgary's audited statement
        # and preserve every underlying Alberta FIR code in provenance.
        derived_sum(
            slug, "net_taxes", d, row_d,
            ["01720","01730","01740","01750","01760","01770","01840"],
            "Alberta LGPI grouping: net municipal property/business/special/local-improvement taxes plus franchise/concession revenue.",
        ),
        derived_sum(
            slug, "user_charges", d, row_d,
            ["01790","01800","01860"],
            "Alberta LGPI grouping: sales to other governments, sales/user charges and rentals.",
        ),
        direct(slug, "investment_income", d, row_d, "01850"),
        derived_sum(
            slug, "developer_contributions", d, row_d,
            ["01885","01960","01962"],
            "Alberta LGPI grouping: contributed/donated assets, developer agreements and offsite levies.",
        ),
        derived_sum(
            slug, "federal_grants", d, row_d,
            ["01892","01902"],
            "Federal operating plus capital transfers.",
        ),
        derived_sum(
            slug, "provincial_grants", d, row_d,
            ["01912","01922"],
            "Provincial operating plus capital transfers.",
        ),
        derived_sum(
            slug, "government_grants_total", d, row_d,
            ["01892","01902","01912","01922","01931","01932"],
            "Federal, provincial and local-government operating and capital transfers.",
        ),
        direct(slug, "total_revenue", d, row_d, "01980"),

        # Expenditure by object.
        derived_sum(
            slug, "grants_object", d, row_d,
            ["02050","02060","02070"],
            "Transfers to governments, local boards/agencies, individuals and organizations.",
        ),
        derived_sum(
            slug, "interest_expense", d, row_d,
            ["02080","02090","02100"],
            "Bank/short-term interest plus operating and capital long-term debt interest.",
        ),
        direct(slug, "salaries_benefits", d, row_d, "02000"),
        direct(slug, "contracted_services", d, row_d, "02010"),
        direct(slug, "goods", d, row_d, "02030"),
        direct(
            slug, "goods_services_other", d, row_d, "02020",
            "Purchases from other governments retained separately inside the LGPI goods/services family.",
        ),
        derived_sum(
            slug, "goods_services_total", d, row_d,
            ["02010","02020","02030"],
            "Contracted/general services, purchases from other governments, and materials/goods/supplies/utilities.",
        ),
        derived_sum(
            slug, "object_other", d, row_d,
            ["02040","02105","02125","02127","02130"],
            "Allowances, asset-retirement accretion, tangible-capital-asset losses/write-downs and other expenditures.",
        ),
        direct(slug, "depreciation_object", d, row_d, "02110"),
        direct(slug, "total_expenditure_object", d, row_d, "02140"),

        # Function/service expenditure. Schedule C expense codes begin at 01170.
        derived_sum(
            slug, "general_government_total", c, row_c,
            ["01170","01180","01190"],
            "Council/legislative, general administration and other general government.",
        ),
        direct(slug, "democracy_costs", c, row_c, "01170"),
        derived_sum(
            slug, "general_government", c, row_c,
            ["01180","01190"],
            "General administration plus other general government, excluding council/legislative costs.",
        ),
        derived_sum(
            slug, "public_safety_total", c, row_c,
            ["01210","01220","01230","01240","01250","01260"],
            "Police, fire, emergency measures, ambulance/first aid, bylaw enforcement and other protective services.",
        ),
        direct(slug, "police", c, row_c, "01210"),
        direct(slug, "fire", c, row_c, "01220"),
        derived_sum(
            slug, "public_safety_other", c, row_c,
            ["01230","01240","01250","01260"],
            "Emergency measures, ambulance/first aid, bylaw enforcement and other protective services.",
        ),
        derived_sum(
            slug, "transportation_total", c, row_c,
            ["01280","01290","01300","01310","01320","01330"],
            "Equipment pool, roads/streets, airport, public transit, storm drainage and other transportation.",
        ),
        direct(slug, "transit", c, row_c, "01310"),
        derived_sum(
            slug, "transportation", c, row_c,
            ["01280","01290","01300","01320","01330"],
            "Transportation functions other than public transit.",
        ),
        derived_sum(
            slug, "environmental_services", c, row_c,
            ["01350","01360","01370","01380"],
            "Water, wastewater, waste management and other environmental protection.",
        ),
        direct(slug, "solid_waste", c, row_c, "01370"),
        derived_sum(
            slug, "social_family_services", c, row_c,
            ["01400","01410"],
            "Family/community support plus day care.",
        ),
        derived_sum(
            slug, "health_services", c, row_c,
            ["01420","01430"],
            "Cemeteries/crematoriums plus other public health and welfare.",
        ),
        direct(slug, "social_housing", c, row_c, "01480"),
        derived_sum(
            slug, "social_services_total", c, row_c,
            ["01400","01410","01420","01430","01480"],
            "Family/community support, day care, public-health/welfare functions and public housing.",
        ),
        derived_sum(
            slug, "planning_development", c, row_c,
            ["01450","01460","01470","01490","01500"],
            "Land-use planning, economic/agricultural development, subdivision development, land/building rentals and other planning/development.",
        ),
        derived_sum(
            slug, "recreation_culture", c, row_c,
            ["01520","01530","01540","01550","01560"],
            "Recreation boards, parks/recreation, culture, convention centres and other recreation/culture.",
        ),
        derived_sum(
            slug, "utility_operations", c, row_c,
            ["01566","01567","01568"],
            "Gas, electric and other utility functions.",
        ),
        direct(slug, "expenditure_other", c, row_c, "01570"),
        direct(slug, "total_expenditure", c, row_c, "01580"),
        direct(
            slug, "depreciation_function", d, row_d, "02110",
            "Alberta FIR reports amortization by object; the same total is surfaced in the legacy LGPI depreciation field.",
        ),
    ]

    # Revenue "Other" is the exhaustive residual after the legacy headline
    # categories. This avoids dropping fines, permits, gains/losses or other
    # jurisdiction-specific revenue lines.
    values = {item["metric"]: int(round(item["valueThousands"] * 1000)) for item in out}
    known = (
        values["net_taxes"]
        + values["user_charges"]
        + values["investment_income"]
        + values["developer_contributions"]
        + values["government_grants_total"]
    )
    other = values["total_revenue"] - known
    out.append(
        observation(
            slug=slug,
            metric="revenue_other",
            value_dollars=other,
            sheet=d,
            source_codes=["01980"],
            method="derived_residual",
            note="Total revenue less LGPI net taxes, user charges, investment income, developer contributions and grants. Captures fines, licences/permits, gains/losses and other revenue without dropping source categories.",
        )
    )
    return out


def validate_calgary(observations: list[dict]) -> None:
    actual = {
        row["metric"]: int(round(float(row["valueThousands"]) * 1000))
        for row in observations
    }
    failures = []
    for metric, expected in CALGARY_CONTROLS.items():
        got = actual.get(metric)
        if got != expected:
            failures.append((metric, expected, got))
    if failures:
        detail = "\n".join(f"{metric}: expected {expected}, got {got}" for metric, expected, got in failures)
        raise RuntimeError("Calgary control reconciliation failed:\n" + detail)


def main() -> None:
    workbook = load_workbook(io.BytesIO(download()), read_only=True, data_only=True)
    needed = ["A(1)-Total", "C(1)-Revenue", "D(1)-Total", "POPL(1)-Population", "ST(1)-Stat"]
    sheets = {name: load_sheet(workbook, name) for name in needed}
    target_codes = find_target_codes(sheets["A(1)-Total"])

    municipalities = []
    all_observations: list[dict] = []
    for slug, code in target_codes.items():
        a_row = sheets["A(1)-Total"].rows[code]
        population_row = sheets["POPL(1)-Population"].rows.get(code, {})
        stats_row = sheets["ST(1)-Stat"].rows.get(code, {})
        observations = build_observations(slug, code, sheets)
        if slug == "calgary":
            validate_calgary(observations)
        all_observations.extend(observations)
        municipalities.append(
            {
                "slug": slug,
                "sourceCode": code,
                "sourceMunicipality": a_row.get("_municipality"),
                "sourceStatus": a_row.get("_status"),
                "population": number(population_row.get("POPL")) or None,
                "dwellingUnitsCandidate": number(stats_row.get("05595")) or None,
                "dwellingUnitsCandidateNote": "Alberta SIR value retained for research only; LGPI household normalization uses Statistics Canada dwellings and this field is not used in calculations.",
            }
        )

    if len(municipalities) != len(TARGETS):
        raise RuntimeError(f"Expected {len(TARGETS)} LGPI Alberta municipalities, emitted {len(municipalities)}")

    document = {
        "schemaVersion": 1,
        "province": "AB",
        "year": 2023,
        "source": {
            "publisher": "Government of Alberta, Municipal Affairs",
            "title": "Municipal Financial and Statistical Data — 2023 financial year",
            "datasetUrl": SOURCE_DATASET_URL,
            "resourceUrl": SOURCE_URL,
            "format": "XLSX",
        },
        "mapping": {
            "name": "LGPI Alberta FIR 2023",
            "version": 1,
            "notes": [
                "Values are stored in thousands of Canadian dollars to match the legacy LGPI display convention.",
                "Direct mappings preserve a single Alberta FIR field code; derived mappings list every contributing code.",
                "Calgary headline values are reconciled against the City of Calgary 2023 audited Annual Financial Report.",
                "Alberta SIR dwelling counts are retained as research context only and are not used for per-household calculations.",
            ],
        },
        "municipalities": municipalities,
        "observations": all_observations,
        "stats": {
            "municipalityCount": len(municipalities),
            "observationCount": len(all_observations),
            "reportedCount": sum(1 for row in all_observations if row["status"] == "reported"),
        },
    }

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(document, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(f"Wrote {OUTPUT}: {len(municipalities)} municipalities, {len(all_observations)} observations")
    for municipality in municipalities:
        print(
            f"- {municipality['slug']}: source={municipality['sourceMunicipality']!r} "
            f"code={municipality['sourceCode']} population={municipality['population']} "
            f"dwelling_candidate={municipality['dwellingUnitsCandidate']}"
        )


if __name__ == "__main__":
    main()

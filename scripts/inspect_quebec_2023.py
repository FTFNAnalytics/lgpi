#!/usr/bin/env python3
"""Inspect Quebec MAMH 2023 open financial data and recover LGPI Quebec universe."""

from __future__ import annotations

import io
import json
import re
import urllib.parse
import urllib.request
from pathlib import Path

from bs4 import BeautifulSoup
from openpyxl import load_workbook

CKAN_ENDPOINTS = [
    "https://www.donneesquebec.ca/recherche/api/3/action/package_show?id=rapport-financier-des-organismes-municipaux-et-autres-documents",
    "https://donneesquebec.ca/recherche/api/3/action/package_show?id=rapport-financier-des-organismes-municipaux-et-autres-documents",
    "https://donnees.iriu.ca/api/3/action/package_show?id=rapport-financier-des-organismes-municipaux-et-autres-documents",
]
LGPI_URL = "https://lgpi.ca/"


def fetch(url: str) -> bytes:
    parsed=urllib.parse.urlsplit(url)
    safe_url=urllib.parse.urlunsplit((
        parsed.scheme,
        parsed.netloc,
        urllib.parse.quote(urllib.parse.unquote(parsed.path), safe="/"),
        parsed.query,
        parsed.fragment,
    ))
    request = urllib.request.Request(
        safe_url,
        headers={"User-Agent":"LGPI-reconstruction/0.1 (+https://github.com/FTFNAnalytics/lgpi)"},
    )
    with urllib.request.urlopen(request, timeout=120) as response:
        return response.read()


def get_package() -> dict:
    failures=[]
    for url in CKAN_ENDPOINTS:
        try:
            data=json.loads(fetch(url))
            if data.get("success") and data.get("result"):
                print(f"CKAN endpoint={url}")
                return data["result"]
            failures.append(f"{url}: success={data.get('success')}")
        except Exception as exc:
            failures.append(f"{url}: {type(exc).__name__}: {exc}")
    raise RuntimeError("No CKAN endpoint succeeded:\n" + "\n".join(failures))


def recover_lgpi_city_options() -> list[str]:
    html=fetch(LGPI_URL)
    soup=BeautifulSoup(html,"html.parser")
    print(f"LGPI homepage bytes={len(html)} selects={len(soup.find_all('select'))}")
    for idx,select in enumerate(soup.find_all("select"),start=1):
        opts=[" ".join(o.stripped_strings) for o in select.find_all("option")]
        print(f"SELECT {idx}: attrs={dict(select.attrs)} option_count={len(opts)} sample={opts[:12]}")

    # Prefer the select with the largest set of non-numeric labels; it is the city selector.
    candidates=[]
    for select in soup.find_all("select"):
        opts=[" ".join(o.stripped_strings).strip() for o in select.find_all("option")]
        labels=[x for x in opts if x and not re.fullmatch(r"\d{4}",x)]
        candidates.append((len(labels),labels))
    if not candidates:
        raise RuntimeError("No select options found on LGPI homepage")
    labels=max(candidates,key=lambda x:x[0])[1]
    # remove prompt labels
    labels=[
        x for x in labels
        if x.lower() not in {"select","browse cities","city","choose a city","- choose a city -"}
    ]
    return labels


def loaded_non_quebec_names() -> set[str]:
    text=Path("lib/lgpi-data.ts").read_text(encoding="utf-8")
    # transparencyRaw tuples: ["slug","Name","PC",...
    return set(re.findall(r'\["[^"]+","([^"]+)","[A-Z]{2}",',text))


def inspect_resource(resource: dict) -> None:
    name=resource.get("name") or resource.get("title") or ""
    url=resource.get("url") or ""
    fmt=(resource.get("format") or "").upper()
    print(f"RESOURCE name={name!r} format={fmt!r} id={resource.get('id')} url={url}")
    if not url:
        return

    low=name.lower()
    if "2023" not in low:
        return
    if fmt == "XLSX" and ("données réelles" in low or "donnees reelles" in low or "financial" in low):
        payload=fetch(url)
        print(f"  XLSX bytes={len(payload)}")
        wb=load_workbook(io.BytesIO(payload),read_only=True,data_only=True)
        print(f"  sheets={wb.sheetnames}")
        for sheet_name in wb.sheetnames[:12]:
            ws=wb[sheet_name]
            print(f"  SHEET {sheet_name!r} rows={ws.max_row} cols={ws.max_column}")
            shown=0
            for row_num,row in enumerate(ws.iter_rows(values_only=True),start=1):
                vals=[str(v)[:160] if v is not None else None for v in row[:16]]
                if any(v not in (None,"") for v in vals):
                    print(f"    ROW {row_num}: {vals}")
                    shown+=1
                if shown>=6:
                    break
    elif fmt == "CSV" and ("description" in low or "simple occurrence" in low or "simple occurrence" in low):
        payload=fetch(url)
        print(f"  CSV bytes={len(payload)} prefix={payload[:500]!r}")


def main() -> None:
    cities=recover_lgpi_city_options()
    loaded=loaded_non_quebec_names()
    missing=sorted({c for c in cities if c not in loaded})
    print(f"LGPI city labels={len(cities)} loaded_names={len(loaded)} missing_from_transparency={len(missing)}")
    print("MISSING CITY LABELS:")
    for city in missing:
        print(f"  - {city}")

    package=get_package()
    resources=package.get("resources") or []
    matches=[]
    for resource in resources:
        name=(resource.get("name") or resource.get("title") or "")
        if "2023" in name and (
            "Données réelles" in name
            or "Donnees reelles" in name
            or "Description poste" in name
            or "Description des postes" in name
            or "financial" in name.lower()
        ):
            matches.append(resource)
    print(f"2023 relevant resources={len(matches)}")
    for resource in matches:
        inspect_resource(resource)


if __name__=="__main__":
    main()
